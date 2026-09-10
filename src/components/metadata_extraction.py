from __future__ import annotations

import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from dotenv import load_dotenv
from huggingface_hub import InferenceClient
from langsmith import traceable
from pydantic import BaseModel, ConfigDict, ValidationError
from pypdf import PdfReader


PROJECT_ROOT = Path(__file__).resolve().parents[2]
ENV_FILE = PROJECT_ROOT / ".env"

load_dotenv(dotenv_path=ENV_FILE, override=False)

DOCUMENTS_DIR = PROJECT_ROOT / "documents"
METADATA_DIR = PROJECT_ROOT / "metadata"
OUTPUT_FILE = METADATA_DIR / "documents.json"
PROMPT_FILE = PROJECT_ROOT / "prompts" / "metadata_extraction.json"

MODEL_NAME = "Qwen/Qwen3-8B"
PROVIDER = "nscale"

MAX_FILE_SIZE_MB = 50
MAX_TEXT_CHARS = 30000
MAX_OUTPUT_TOKENS = 2500
RETRY_MAX_TOKENS = 1800

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".docx",
    ".json",
    ".csv",
}


class DocumentMetadata(BaseModel):
    model_config = ConfigDict(extra="ignore")

    title: str
    summary: str
    document_type: str
    topics: list[str]
    keywords: list[str]
    people: list[str]
    organizations: list[str]
    locations: list[str]
    dates: list[str]
    language: str
    department: str
    author: str
    version: str
    entities: list[str]


class DocumentMetadataRecord(BaseModel):
    document_id: str
    filename: str
    relative_path: str
    extension: str
    size_bytes: int
    modified_at: str
    extracted_metadata: DocumentMetadata


class MetadataExtractionError(RuntimeError):
    pass


class MetadataExtractor:
    def __init__(
        self,
        model_name: str = MODEL_NAME,
        provider: str = PROVIDER,
        prompt_file: str | Path = PROMPT_FILE,
    ) -> None:
        self.model_name = model_name
        self.provider = provider
        self.prompt_file = Path(prompt_file)

        token = os.getenv("HF_TOKEN")

        if not token:
            raise EnvironmentError(
                f"HF_TOKEN is not set. Expected it in: {ENV_FILE}"
            )

        self.system_prompt = self._load_prompt()

        self.client = InferenceClient(
            provider=self.provider,
            api_key=token,
        )

    def _load_prompt(self) -> str:
        if not self.prompt_file.exists():
            raise FileNotFoundError(
                f"Prompt file not found: {self.prompt_file}"
            )

        try:
            data = json.loads(
                self.prompt_file.read_text(
                    encoding="utf-8"
                )
            )
        except json.JSONDecodeError as exc:
            raise ValueError(
                f"Invalid JSON in prompt file: {self.prompt_file}"
            ) from exc

        if not isinstance(data, dict):
            raise ValueError(
                "Prompt JSON must contain an object."
            )

        direct_prompt = data.get("system")

        if isinstance(direct_prompt, str):
            return direct_prompt

        section = data.get("metadata_extraction")

        if isinstance(section, dict):
            prompt = section.get("system")

            if isinstance(prompt, str):
                return prompt

        raise ValueError(
            "Prompt JSON must contain either "
            "'system' or 'metadata_extraction.system'."
        )

    @staticmethod
    def _schema() -> dict[str, Any]:
        schema = DocumentMetadata.model_json_schema()

        schema["required"] = list(
            schema["properties"].keys()
        )
        schema["additionalProperties"] = False

        return {
            "name": "document_metadata",
            "description": "Metadata extracted from a document for RAG filtering.",
            "schema": schema,
            "strict": True,
        }

    def _build_messages(
        self,
        filename: str,
        text: str,
    ) -> list[dict[str, str]]:
        user_prompt = (
            "Extract document metadata using the supplied schema.\n\n"
            f"Filename:\n{filename}\n\n"
            f"Document content:\n{text}\n\n"
            "Every field must be present in the JSON response. "
            "Use an empty string for an unknown scalar field and "
            "an empty array for an unknown list field. "
            "Do not invent or infer unsupported information. "
            "Return only JSON."
        )

        return [
            {
                "role": "system",
                "content": self.system_prompt,
            },
            {
                "role": "user",
                "content": user_prompt,
            },
        ]

    @traceable(
        name="document_metadata_extraction",
        run_type="llm",
        tags=[
            "rag",
            "metadata-extraction",
            "huggingface",
            "qwen3-8b",
            "nscale",
        ],
    )
    def extract(
        self,
        filename: str,
        text: str,
    ) -> DocumentMetadata:
        if not text.strip():
            return DocumentMetadata(
                title="",
                summary="",
                document_type="",
                topics=[],
                keywords=[],
                people=[],
                organizations=[],
                locations=[],
                dates=[],
                language="",
                department="",
                author="",
                version="",
                entities=[],
            )

        attempts = [
            (
                text,
                MAX_OUTPUT_TOKENS,
                True,
            ),
            (
                compact_text_for_retry(text),
                RETRY_MAX_TOKENS,
                False,
            ),
        ]

        last_error: Exception | None = None

        for attempt_number, (
            candidate_text,
            max_tokens,
            strict_schema,
        ) in enumerate(attempts, start=1):
            try:
                kwargs: dict[str, Any] = {
                    "model": self.model_name,
                    "messages": self._build_messages(
                        filename=filename,
                        text=candidate_text,
                    ),
                    "temperature": 0.7,
                    "top_p": 0.8,
                    "max_tokens": max_tokens,
                    "extra_body": {
                        "chat_template_kwargs": {
                            "enable_thinking": False
                        }
                    },
                }

                if strict_schema:
                    kwargs["response_format"] = {
                        "type": "json_schema",
                        "json_schema": self._schema(),
                    }
                else:
                    kwargs["response_format"] = {
                        "type": "json_object"
                    }

                response = (
                    self.client
                    .chat
                    .completions
                    .create(**kwargs)
                )

                content = self._extract_content(response)

                if not content:
                    raise MetadataExtractionError(
                        f"Model returned an empty response on attempt "
                        f"{attempt_number}."
                    )

                content = clean_json_text(content)

                try:
                    payload = json.loads(content)
                except json.JSONDecodeError as exc:
                    raise MetadataExtractionError(
                        f"Model returned invalid JSON on attempt "
                        f"{attempt_number}: {content}"
                    ) from exc

                try:
                    metadata = DocumentMetadata.model_validate(
                        payload
                    )
                except ValidationError as exc:
                    raise MetadataExtractionError(
                        f"Metadata validation failed on attempt "
                        f"{attempt_number}: {exc}"
                    ) from exc

                return sanitize_metadata(metadata)

            except Exception as exc:
                last_error = exc

                if attempt_number < len(attempts):
                    continue

        raise MetadataExtractionError(
            f"Metadata extraction failed after "
            f"{len(attempts)} attempts: {last_error}"
        )

    @staticmethod
    def _extract_content(response: Any) -> str:
        if not getattr(response, "choices", None):
            raise MetadataExtractionError(
                "Model returned no completion choices."
            )

        message = response.choices[0].message
        content = message.content

        if isinstance(content, str):
            return content.strip()

        if isinstance(content, list):
            parts: list[str] = []

            for item in content:
                if isinstance(item, dict):
                    value = item.get("text")
                    if isinstance(value, str):
                        parts.append(value)

            return "".join(parts).strip()

        return ""

    def extract_from_file(
        self,
        path: Path,
    ) -> DocumentMetadata:
        text = read_document(path)

        prepared_text = prepare_text(
            text,
            MAX_TEXT_CHARS,
        )

        return self.extract(
            filename=path.name,
            text=prepared_text,
        )


def sanitize_metadata(
    metadata: DocumentMetadata,
) -> DocumentMetadata:
    metadata.title = clean_scalar(metadata.title)
    metadata.summary = clean_scalar(metadata.summary)
    metadata.document_type = clean_scalar(
        metadata.document_type
    )
    metadata.language = clean_scalar(
        metadata.language
    )
    metadata.department = clean_scalar(
        metadata.department
    )
    metadata.author = clean_scalar(
        metadata.author
    )
    metadata.version = sanitize_version(
        metadata.version
    )

    list_fields = [
        "topics",
        "keywords",
        "people",
        "organizations",
        "locations",
        "dates",
        "entities",
    ]

    for field_name in list_fields:
        values = getattr(metadata, field_name)

        cleaned: list[str] = []
        seen: set[str] = set()

        for value in values:
            if not isinstance(value, str):
                continue

            value = value.strip()

            if not value:
                continue

            key = value.casefold()

            if key in seen:
                continue

            seen.add(key)
            cleaned.append(value)

        setattr(
            metadata,
            field_name,
            cleaned,
        )

    return metadata


def clean_scalar(value: str) -> str:
    if not isinstance(value, str):
        return ""

    value = value.strip()

    invalid_values = {
        "",
        "unknown",
        "unspecified",
        "not specified",
        "not available",
        "n/a",
        "na",
        "none",
        "null",
        "not mentioned",
        "not stated",
    }

    if value.casefold() in invalid_values:
        return ""

    return value


def sanitize_version(value: str) -> str:
    value = clean_scalar(value)

    if not value:
        return ""

    lowered = value.casefold()

    invalid_fragments = {
        "assuming",
        "default version",
        "based on document structure",
        "unspecified in document",
        "not explicitly",
        "not stated",
        "not specified",
        "unknown version",
    }

    if any(
        fragment in lowered
        for fragment in invalid_fragments
    ):
        return ""

    if len(value) > 50:
        return ""

    if not re.search(r"\d", value):
        return ""

    if re.fullmatch(
        r"(?:\d+\.){4,}\d+",
        value,
    ):
        return ""

    return value


def clean_json_text(content: str) -> str:
    content = content.strip()

    if content.startswith("```"):
        content = re.sub(
            r"^```(?:json)?\s*",
            "",
            content,
            flags=re.IGNORECASE,
        )
        content = re.sub(
            r"\s*```$",
            "",
            content,
        )

    start = content.find("{")
    end = content.rfind("}")

    if start >= 0 and end >= start:
        content = content[start:end + 1]

    return content.strip()


def compact_text_for_retry(text: str) -> str:
    text = normalize_text(text)

    if len(text) <= 12000:
        return text

    first = 5000
    middle_start = len(text) // 2 - 2500
    middle_start = max(0, middle_start)
    last = 5000

    return (
        text[:first]
        + "\n\n[...MIDDLE SAMPLE...]\n\n"
        + text[middle_start:middle_start + 5000]
        + "\n\n[...END SAMPLE...]\n\n"
        + text[-last:]
    )


def normalize_text(text: str) -> str:
    return "\n".join(
        line.rstrip()
        for line in text.replace(
            "\x00",
            "",
        ).splitlines()
    ).strip()


def read_document(path: Path) -> str:
    extension = path.suffix.lower()

    if extension in {
        ".txt",
        ".md",
        ".markdown",
        ".csv",
    }:
        return path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    if extension == ".json":
        raw = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        try:
            return json.dumps(
                json.loads(raw),
                ensure_ascii=False,
                indent=2,
            )
        except json.JSONDecodeError:
            return raw

    if extension == ".pdf":
        reader = PdfReader(str(path))
        pages: list[str] = []

        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""
            except Exception:
                page_text = ""

            if page_text.strip():
                pages.append(page_text)

        return "\n\n".join(pages)

    if extension == ".docx":
        document = Document(str(path))
        parts: list[str] = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if value:
                parts.append(value)

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(values):
                    parts.append(
                        " | ".join(values)
                    )

        return "\n".join(parts)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


def prepare_text(
    text: str,
    max_chars: int,
) -> str:
    text = normalize_text(text)

    if not text:
        return ""

    if len(text) <= max_chars:
        return text

    first = int(max_chars * 0.50)
    last = int(max_chars * 0.25)
    middle = max_chars - first - last

    middle_start = max(
        0,
        (len(text) // 2) - (middle // 2),
    )

    return (
        text[:first]
        + "\n\n[...MIDDLE CONTENT...]\n\n"
        + text[
            middle_start:
            middle_start + middle
        ]
        + "\n\n[...FINAL CONTENT...]\n\n"
        + text[-last:]
    )


def create_document_id(
    path: Path,
) -> str:
    relative_path = (
        path
        .relative_to(PROJECT_ROOT)
        .as_posix()
    )

    return hashlib.sha256(
        relative_path.encode("utf-8")
    ).hexdigest()[:16]


def get_modified_at(
    path: Path,
) -> str:
    return datetime.fromtimestamp(
        path.stat().st_mtime,
        tz=timezone.utc,
    ).isoformat()


def discover_documents() -> list[Path]:
    if not DOCUMENTS_DIR.exists():
        raise FileNotFoundError(
            f"Documents directory not found: "
            f"{DOCUMENTS_DIR}"
        )

    documents: list[Path] = []

    for path in DOCUMENTS_DIR.rglob("*"):
        if not path.is_file():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        if (
            path.stat().st_size
            > MAX_FILE_SIZE_MB * 1024 * 1024
        ):
            print(
                f"Skipping oversized file: {path}"
            )
            continue

        documents.append(path)

    return sorted(
        documents,
        key=lambda path: path.as_posix().lower(),
    )


def load_existing_records() -> dict[str, Any]:
    if not OUTPUT_FILE.exists():
        return {}

    try:
        payload = json.loads(
            OUTPUT_FILE.read_text(
                encoding="utf-8"
            )
        )
    except json.JSONDecodeError:
        return {}

    records = payload.get(
        "documents",
        {}
    )

    return (
        records
        if isinstance(records, dict)
        else {}
    )


def needs_processing(
    path: Path,
    existing: dict[str, Any],
) -> bool:
    document_id = create_document_id(path)
    record = existing.get(document_id)

    if not record:
        return True

    if record.get("error"):
        return True

    if (
        record.get("modified_at")
        != get_modified_at(path)
    ):
        return True

    if (
        record.get("size_bytes")
        != path.stat().st_size
    ):
        return True

    return False


def build_record(
    path: Path,
    metadata: DocumentMetadata,
) -> DocumentMetadataRecord:
    return DocumentMetadataRecord(
        document_id=create_document_id(path),
        filename=path.name,
        relative_path=(
            path
            .relative_to(PROJECT_ROOT)
            .as_posix()
        ),
        extension=path.suffix.lower(),
        size_bytes=path.stat().st_size,
        modified_at=get_modified_at(path),
        extracted_metadata=metadata,
    )


def write_output(
    records: dict[str, Any],
) -> None:
    METADATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    payload = {
        "generated_at":
            datetime.now(
                timezone.utc
            ).isoformat(),
        "document_count":
            len(records),
        "documents":
            records,
    }

    OUTPUT_FILE.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def main() -> None:
    extractor = MetadataExtractor()

    documents = discover_documents()
    existing = load_existing_records()
    records = dict(existing)

    print(
        f"Project root: {PROJECT_ROOT}"
    )
    print(
        f"Documents discovered: {len(documents)}"
    )
    print(
        f"Model: {MODEL_NAME}"
    )
    print(
        f"Provider: {PROVIDER}"
    )
    print()

    processed = 0
    skipped = 0
    failed = 0

    for index, path in enumerate(
        documents,
        start=1,
    ):
        relative_path = path.relative_to(
            PROJECT_ROOT
        )

        if not needs_processing(
            path,
            existing,
        ):
            print(
                f"[{index}/{len(documents)}] "
                f"SKIP    {relative_path}"
            )
            skipped += 1
            continue

        print(
            f"[{index}/{len(documents)}] "
            f"PROCESS {relative_path}"
        )

        try:
            metadata = extractor.extract_from_file(
                path
            )

            record = build_record(
                path,
                metadata,
            )

            records[
                record.document_id
            ] = record.model_dump()

            processed += 1

            print(
                f"[{index}/{len(documents)}] "
                f"SUCCESS {relative_path}"
            )

        except Exception as exc:
            document_id = create_document_id(path)

            records[document_id] = {
                "document_id": document_id,
                "filename": path.name,
                "relative_path": relative_path.as_posix(),
                "extension": path.suffix.lower(),
                "size_bytes": path.stat().st_size,
                "modified_at": get_modified_at(path),
                "error": str(exc),
            }

            failed += 1

            print(
                f"[{index}/{len(documents)}] "
                f"FAILED  {relative_path}"
            )
            print(
                f"Reason: {exc}"
            )

    current_ids = {
        create_document_id(path)
        for path in documents
    }

    records = {
        document_id: record
        for document_id, record in records.items()
        if document_id in current_ids
    }

    write_output(records)

    print()
    print(
        f"Processed: {processed}"
    )
    print(
        f"Skipped:   {skipped}"
    )
    print(
        f"Failed:    {failed}"
    )
    print()
    print(
        f"Metadata written to: "
        f"{OUTPUT_FILE}"
    )


if __name__ == "__main__":
    main()
